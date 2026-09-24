"""Document parsing: bytes plus a filename in, pages of text out.

Two rules shape this module.

**A failure is never an empty success.** The tempting implementation wraps the
whole body in ``except Exception: return []`` and moves on. That converts a
corrupt upload, an unsupported format and a genuinely empty file into the same
outcome — a document that ingests "successfully" and can never be retrieved —
and the operator has no signal. Every failure here raises, and a page that
parses to nothing is reported as a warning rather than dropped, because "the
text extractor found nothing on page 7" is information a user needs.

**Validation happens before parsing.** Extension and size are checked first, so
an unsupported or oversized upload is rejected without spending CPU on it and
without a partially-registered document.
"""

from __future__ import annotations

import io
from pathlib import Path

from pydantic import BaseModel, Field

from coursellm.core import config as config_module
from coursellm.core.config import Settings
from coursellm.core.errors import (
    PayloadTooLargeError,
    UnsupportedMediaTypeError,
    ValidationError,
)
from coursellm.db.models.content import SourceType

# Filename fragments that identify provenance. Order matters: the first match
# wins, and the order runs from most specific to most generic so that
# ``lecture-slides.pdf`` is classified as a slide deck rather than a lecture.
_SOURCE_TYPE_HINTS: tuple[tuple[tuple[str, ...], SourceType], ...] = (
    (("syllabus", "syllabi", "curriculum"), SourceType.SYLLABUS),
    (("slide", "slides", "deck", "keynote"), SourceType.SLIDE),
    (("paper", "arxiv", "journal", "article", "preprint"), SourceType.PAPER),
    (("book", "chapter", "textbook", "monograph"), SourceType.BOOK),
    (("lecture", "lect", "class", "tutorial", "seminar"), SourceType.LECTURE),
    (("note", "notes", "summary", "cheatsheet", "cheat-sheet"), SourceType.NOTES),
    (
        ("doc", "docs", "documentation", "manual", "readme", "guide", "reference", "spec"),
        SourceType.DOCUMENTATION,
    ),
)


class ParsedPage(BaseModel):
    """One page (or, for flat formats, one logical page) of extracted text."""

    page: int = Field(ge=1, description="1-based page number, as cited to the user.")
    text: str = Field(description="Extracted text. May be empty for a scanned page.")


class ParsedDocument(BaseModel):
    """The result of parsing one upload.

    ``warnings`` is part of the contract rather than a log line: it is surfaced
    to the uploader and stored on the ingestion record, so a document that
    ingested with three unreadable pages is visibly degraded rather than
    silently incomplete.
    """

    pages: list[ParsedPage]
    page_count: int
    warnings: list[str] = Field(default_factory=list)


def _resolve_settings(settings: Settings | None) -> Settings:
    """Use an explicit settings object when given, else the process settings.

    Read through the module rather than importing ``settings`` directly so that
    the test fixture which replaces the accessor is honoured.
    """
    return settings if settings is not None else config_module.get_settings()


def _extension(filename: str) -> str:
    return Path(filename).suffix.lower()


def sniff_source_type(filename: str) -> SourceType:
    """Classify provenance from the filename.

    A heuristic, not a classifier: it exists so that the common case (a lecturer
    names the file ``lecture-4-transformers.pdf``) is labelled without a model
    call. The default is :attr:`SourceType.OTHER` rather than a guess, because a
    wrong trust label is worse than an absent one.
    """
    lowered = Path(filename).name.lower()
    for fragments, source_type in _SOURCE_TYPE_HINTS:
        if any(fragment in lowered for fragment in fragments):
            return source_type
    return SourceType.OTHER


def parse_document(
    data: bytes,
    *,
    filename: str,
    content_type: str | None,
    settings: Settings | None = None,
) -> ParsedDocument:
    """Parse ``data`` into pages of text.

    ``settings`` is optional so that a caller which already holds the resolved
    :class:`~coursellm.core.config.Settings` (the ingestion pipeline) validates
    against exactly the limits it was configured with, while a quick script can
    omit it. It never falls back to a default silently: the ambient settings are
    read if it is absent.
    """
    resolved = _resolve_settings(settings)

    extension = _extension(filename)
    if not extension:
        raise UnsupportedMediaTypeError(
            f"{filename!r} has no file extension; cannot determine how to parse it."
        )
    if extension not in resolved.allowed_extensions:
        allowed = ", ".join(sorted(resolved.allowed_extensions))
        raise UnsupportedMediaTypeError(
            f"{extension!r} files are not accepted (allowed: {allowed})."
        )

    # Checked before parsing so an oversized upload costs nothing beyond the
    # bytes already buffered by the caller.
    if len(data) > resolved.max_upload_bytes:
        raise PayloadTooLargeError(
            f"File is {len(data)} bytes; the limit is {resolved.max_upload_bytes} bytes."
        )

    if extension == ".pdf":
        return _parse_pdf(data)
    if extension in {".txt", ".md", ".markdown"}:
        return _parse_plain_text(data, extension)
    if extension == ".docx":
        return _parse_docx(data)
    raise UnsupportedMediaTypeError(
        f"Parsing {extension!r} files (content type {content_type or 'unknown'!r}) "
        "is not implemented yet."
    )


def _parse_plain_text(data: bytes, extension: str) -> ParsedDocument:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValidationError(
            f"The {extension} file is not valid UTF-8 text; re-encode it and upload again."
        ) from exc

    warnings: list[str] = []
    if not text.strip():
        warnings.append("The file contained no text.")
    return ParsedDocument(
        pages=[ParsedPage(page=1, text=text)],
        page_count=1,
        warnings=warnings,
    )


def _parse_pdf(data: bytes) -> ParsedDocument:
    """Extract text page by page with ``pypdf``.

    Scanned PDFs legitimately produce empty pages; those become warnings. A
    structural failure (truncated file, not a PDF at all) raises, because
    returning an empty document there would permanently poison the dedup key for
    those bytes.
    """
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        page_count = len(reader.pages)
    except (PdfReadError, ValueError, OSError, TypeError) as exc:
        raise ValidationError(f"Could not read the PDF: {exc}") from exc

    if page_count == 0:
        raise ValidationError("The PDF contains no pages.")

    pages: list[ParsedPage] = []
    warnings: list[str] = []
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            raise ValidationError(
                f"Could not extract text from PDF page {page_number}: {exc}"
            ) from exc
        if not text.strip():
            warnings.append(f"Page {page_number} produced no extractable text.")
        pages.append(ParsedPage(page=page_number, text=text))

    return ParsedDocument(pages=pages, page_count=len(pages), warnings=warnings)


def _parse_docx(data: bytes) -> ParsedDocument:
    """Extract paragraph text with ``python-docx``.

    The dependency is an optional extra, so its absence is reported as an
    unsupported media type with the install hint rather than an ``ImportError``
    that a user cannot act on. DOCX has no pagination in the format itself, so
    the whole document is one logical page.
    """
    try:
        import docx
    except ImportError as exc:
        raise UnsupportedMediaTypeError(
            "DOCX support requires the optional 'docx' dependency (pip install 'coursellm[docx]')."
        ) from exc

    try:
        document = docx.Document(io.BytesIO(data))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    except Exception as exc:
        raise ValidationError(f"Could not read the .docx file: {exc}") from exc

    warnings: list[str] = []
    if not text.strip():
        warnings.append("The file contained no paragraph text.")
    return ParsedDocument(
        pages=[ParsedPage(page=1, text=text)],
        page_count=1,
        warnings=warnings,
    )
