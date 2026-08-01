import fitz  # PyMuPDF
from pptx import Presentation
import os
import base64
from groq import Groq
from core.config import settings

def parse_document(file_path: str) -> list[dict]:
    """
    Parses a document (PDF, PPTX, TXT, PNG, JPG) and returns a list of dictionaries containing
    text and page/slide numbers.
    Format: [{"page": 1, "text": "Content here..."}, ...]
    """
    _, ext = os.path.splitext(file_path)
    ext = ext.lower()
    
    if ext == ".pdf":
        return _parse_pdf(file_path)
    elif ext == ".pptx":
        return _parse_pptx(file_path)
    elif ext == ".txt":
        return _parse_txt(file_path)
    elif ext in [".png", ".jpg", ".jpeg"]:
        return _parse_image(file_path)
    else:
        raise ValueError(f"Unsupported file extension: {ext}")

def _parse_txt(file_path: str) -> list[dict]:
    pages_data = []
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            text = f.read().strip()
            if text:
                pages_data.append({
                    "page": 1,
                    "text": text
                })
    except Exception as e:
        print(f"Error parsing TXT {file_path}: {e}")
    return pages_data

def _parse_image(file_path: str) -> list[dict]:
    pages_data = []
    try:
        # Encode image to base64
        with open(file_path, "rb") as image_file:
            encoded_string = base64.b64encode(image_file.read()).decode('utf-8')
            
        # Get image mime type
        _, ext = os.path.splitext(file_path)
        ext = ext.lower().replace(".", "")
        mime_type = f"image/jpeg" if ext in ["jpg", "jpeg"] else f"image/{ext}"
        
        # Initialize Groq client
        client = Groq(api_key=settings.GROQ_API_KEY)
        
        prompt = (
            "You are an expert OCR and Computer Vision assistant. "
            "Please perfectly transcribe all text, formulas, and code present in this image. "
            "If there are any charts, diagrams, or visual structures, describe them in deep detail "
            "so that a student could understand the concept without seeing the image. "
            "Return ONLY the transcribed text and descriptions, with no conversational filler."
        )
        
        response = client.chat.completions.create(
            messages=[
                {
                    "role": "user",
                    "content": [
                        {"type": "text", "text": prompt},
                        {
                            "type": "image_url",
                            "image_url": {
                                "url": f"data:{mime_type};base64,{encoded_string}",
                            },
                        },
                    ],
                }
            ],
            model="llama-3.2-90b-vision-preview",
            temperature=0.0
        )
        
        text = response.choices[0].message.content.strip()
        
        if text:
            pages_data.append({
                "page": 1,
                "text": text
            })
            
    except Exception as e:
        print(f"Error parsing Image {file_path}: {e}")
        
    return pages_data

def _parse_pdf(file_path: str) -> list[dict]:
    pages_data = []
    try:
        doc = fitz.open(file_path)
        for i, page in enumerate(doc):
            text = page.get_text("text")
            if text.strip():
                pages_data.append({
                    "page": i + 1,
                    "text": text.strip()
                })
    except Exception as e:
        print(f"Error parsing PDF {file_path}: {e}")
    return pages_data

def _parse_pptx(file_path: str) -> list[dict]:
    slides_data = []
    try:
        prs = Presentation(file_path)
        for i, slide in enumerate(prs.slides):
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    slide_text.append(shape.text.strip())
            
            full_text = "\n".join([t for t in slide_text if t])
            if full_text:
                slides_data.append({
                    "page": i + 1,
                    "text": full_text
                })
    except Exception as e:
        print(f"Error parsing PPTX {file_path}: {e}")
    return slides_data
