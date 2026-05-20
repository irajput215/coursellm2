import fitz  # PyMuPDF
from pptx import Presentation
import os

def parse_document(file_path: str) -> list[dict]:
    """
    Parses a document (PDF or PPTX) and returns a list of dictionaries containing
    text and page/slide numbers.
    Format: [{"page": 1, "text": "Content here..."}, ...]
    """
    _, ext = os.path.splitext(file_path)
    ext = ext.lower()
    
    if ext == ".pdf":
        return _parse_pdf(file_path)
    elif ext == ".pptx":
        return _parse_pptx(file_path)
    else:
        raise ValueError(f"Unsupported file extension: {ext}")

def _parse_pdf(file_path: str) -> list[dict]:
    pages_data = []
    try:
        doc = fitz.open(file_path)
        for i, page in enumerate(doc):
            text = page.get_text("text")
            if text.strip():
                pages_data.append({
                    "page": i + 1,
                    "text": text.strip()
                })
    except Exception as e:
        print(f"Error parsing PDF {file_path}: {e}")
    return pages_data

def _parse_pptx(file_path: str) -> list[dict]:
    slides_data = []
    try:
        prs = Presentation(file_path)
        for i, slide in enumerate(prs.slides):
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    slide_text.append(shape.text.strip())
            
            full_text = "\n".join([t for t in slide_text if t])
            if full_text:
                slides_data.append({
                    "page": i + 1,
                    "text": full_text
                })
    except Exception as e:
        print(f"Error parsing PPTX {file_path}: {e}")
    return slides_data
